# -*- coding: utf-8 -*-
"""
make_report.py
Constructs the submission-ready academic report for SecureJobLab.
Course Code: 20CYS403 (Web Application Security)
Output: C:\\Users\\Ram\\Desktop\\SecureWebLab\\SecureJobLab_Web_Application_Security_Report.docx
"""

import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

COLOR_NAVY = RGBColor(10, 37, 64)       # #0A2540
COLOR_SLATE = RGBColor(30, 41, 59)      # #1E293B
COLOR_BODY = RGBColor(51, 65, 85)       # #334155
COLOR_BLUE = RGBColor(2, 132, 199)      # #0284C7
COLOR_RED = RGBColor(207, 19, 34)       # #CF1322
COLOR_GREEN = RGBColor(22, 163, 74)     # #16A34A
COLOR_MUTED = RGBColor(100, 116, 139)   # #64748B

HEX_NAVY = "0A2540"
HEX_SLATE = "1E293B"
HEX_LIGHT_BLUE = "F0F9FF"
HEX_LIGHT_GRAY = "F8FAFC"
HEX_BORDER = "CBD5E1"
HEX_RED_BG = "FFF1F0"
HEX_RED_BORDER = "FFA39E"
HEX_GREEN_BG = "F6FFED"
HEX_GREEN_BORDER = "B7EB8F"

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    edges = {'top': top, 'bottom': bottom, 'left': left, 'right': right}
    for edge, border_def in edges.items():
        if border_def:
            val = border_def.get('val', 'single')
            sz = border_def.get('sz', '4')
            color = border_def.get('color', 'auto')
            el = parse_xml(f'<w:{edge} {nsdecls("w")} w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>')
            tcBorders.append(el)
        else:
            el = parse_xml(f'<w:{edge} {nsdecls("w")} w:val="none"/>')
            tcBorders.append(el)
    tcPr.append(tcBorders)

def add_header_footer(doc):
    for s in doc.sections:
        s.top_margin = Inches(0.85)
        s.bottom_margin = Inches(0.85)
        s.left_margin = Inches(0.85)
        s.right_margin = Inches(0.85)

        header = s.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("Course: 20CYS403 | Web Application Security Laboratory Report | SecureJobLab")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = COLOR_MUTED

        footer = s.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("SecureJobLab Platform — Exactly 5 CWE Security Modules | Page ")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = COLOR_MUTED
        fldSimple = OxmlElement('w:fldSimple')
        fldSimple.set(qn('w:instr'), 'PAGE')
        fp._p.append(fldSimple)

def add_h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(15)
    run.font.bold = True
    run.font.color.rgb = COLOR_NAVY
    return p

def add_h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(11)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(12.5)
    run.font.bold = True
    run.font.color.rgb = COLOR_SLATE
    return p

def add_h3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = COLOR_BLUE
    return p

def add_p(doc, text, bold_prefix=None, space_after=3):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_SLATE
    r_body = p.add_run(text)
    r_body.font.name = "Calibri"
    r_body.font.size = Pt(10)
    r_body.font.color.rgb = COLOR_BODY
    return p

def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.12
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(9.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_SLATE
    r_body = p.add_run(text)
    r_body.font.name = "Calibri"
    r_body.font.size = Pt(9.5)
    r_body.font.color.rgb = COLOR_BODY
    return p

def add_callout(doc, text, title="SECURITY SUMMARY", ctype="blue"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.5)
    cell = tbl.cell(0, 0)
    set_cell_margins(cell, top=70, bottom=70, left=120, right=120)

    if ctype == "red":
        bg_col, border_col, title_col, bar_col = HEX_RED_BG, HEX_RED_BORDER, COLOR_RED, "CF1322"
    elif ctype == "green":
        bg_col, border_col, title_col, bar_col = HEX_GREEN_BG, HEX_GREEN_BORDER, COLOR_GREEN, "16A34A"
    else:
        bg_col, border_col, title_col, bar_col = HEX_LIGHT_BLUE, "BAE6FD", COLOR_BLUE, "0284C7"

    set_cell_background(cell, bg_col)
    set_cell_borders(cell,
        top={'val': 'single', 'sz': '4', 'color': border_col},
        bottom={'val': 'single', 'sz': '4', 'color': border_col},
        right={'val': 'single', 'sz': '4', 'color': border_col},
        left={'val': 'single', 'sz': '24', 'color': bar_col}
    )

    cp = cell.paragraphs[0]
    cp.paragraph_format.space_before = Pt(0)
    cp.paragraph_format.space_after = Pt(0)
    rt = cp.add_run(f"[{title}] ")
    rt.font.name = "Calibri"
    rt.font.size = Pt(9.5)
    rt.font.bold = True
    rt.font.color.rgb = title_col

    rb = cp.add_run(text)
    rb.font.name = "Calibri"
    rb.font.size = Pt(9)
    rb.font.color.rgb = COLOR_BODY

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(3)

def add_code_block(doc, code_str, caption=None):
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.paragraph_format.space_before = Pt(3)
        p_cap.paragraph_format.space_after = Pt(1)
        rc = p_cap.add_run(f"Listing: {caption}")
        rc.font.name = "Calibri"
        rc.font.size = Pt(8.5)
        rc.font.bold = True
        rc.font.color.rgb = COLOR_MUTED

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.5)
    cell = tbl.cell(0, 0)
    set_cell_margins(cell, top=50, bottom=50, left=90, right=90)
    set_cell_background(cell, HEX_LIGHT_GRAY)
    set_cell_borders(cell,
        top={'val': 'single', 'sz': '4', 'color': HEX_BORDER},
        bottom={'val': 'single', 'sz': '4', 'color': HEX_BORDER},
        right={'val': 'single', 'sz': '4', 'color': HEX_BORDER},
        left={'val': 'single', 'sz': '18', 'color': "0284C7"}
    )
    cp = cell.paragraphs[0]
    cp.paragraph_format.space_before = Pt(0)
    cp.paragraph_format.space_after = Pt(0)
    cp.paragraph_format.line_spacing = 1.05
    run = cp.add_run(code_str)
    run.font.name = "Consolas"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(30, 41, 59)

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(3)

def add_dual_code_comparison(doc, vuln_code, sec_code, vuln_title="Vulnerable Code (Unsafe)", sec_title="Secure Mitigated Code (Defended)"):
    tbl = doc.add_table(rows=2, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(3.2)
    tbl.columns[1].width = Inches(3.2)

    c0, c1 = tbl.cell(0, 0), tbl.cell(0, 1)
    set_cell_background(c0, HEX_RED_BG)
    set_cell_background(c1, HEX_GREEN_BG)
    set_cell_margins(c0, top=40, bottom=40, left=60, right=60)
    set_cell_margins(c1, top=40, bottom=40, left=60, right=60)
    set_cell_borders(c0,
        top={'val': 'single', 'sz': '4', 'color': HEX_RED_BORDER},
        bottom={'val': 'single', 'sz': '4', 'color': HEX_RED_BORDER},
        left={'val': 'single', 'sz': '12', 'color': "CF1322"},
        right={'val': 'single', 'sz': '4', 'color': HEX_RED_BORDER}
    )
    set_cell_borders(c1,
        top={'val': 'single', 'sz': '4', 'color': HEX_GREEN_BORDER},
        bottom={'val': 'single', 'sz': '4', 'color': HEX_GREEN_BORDER},
        left={'val': 'single', 'sz': '12', 'color': "16A34A"},
        right={'val': 'single', 'sz': '4', 'color': HEX_GREEN_BORDER}
    )

    p0 = c0.paragraphs[0]
    p0.paragraph_format.space_before = Pt(0)
    p0.paragraph_format.space_after = Pt(0)
    r0 = p0.add_run(f"🔴 {vuln_title}")
    r0.font.name = "Calibri"
    r0.font.size = Pt(8.5)
    r0.font.bold = True
    r0.font.color.rgb = COLOR_RED

    p1 = c1.paragraphs[0]
    p1.paragraph_format.space_before = Pt(0)
    p1.paragraph_format.space_after = Pt(0)
    r1 = p1.add_run(f"🟢 {sec_title}")
    r1.font.name = "Calibri"
    r1.font.size = Pt(8.5)
    r1.font.bold = True
    r1.font.color.rgb = COLOR_GREEN

    cb0, cb1 = tbl.cell(1, 0), tbl.cell(1, 1)
    set_cell_background(cb0, HEX_LIGHT_GRAY)
    set_cell_background(cb1, HEX_LIGHT_GRAY)
    set_cell_margins(cb0, top=50, bottom=50, left=60, right=60)
    set_cell_margins(cb1, top=50, bottom=50, left=60, right=60)
    set_cell_borders(cb0,
        top={'val': 'none'},
        bottom={'val': 'single', 'sz': '4', 'color': HEX_BORDER},
        left={'val': 'single', 'sz': '12', 'color': "CF1322"},
        right={'val': 'single', 'sz': '4', 'color': HEX_BORDER}
    )
    set_cell_borders(cb1,
        top={'val': 'none'},
        bottom={'val': 'single', 'sz': '4', 'color': HEX_BORDER},
        left={'val': 'single', 'sz': '12', 'color': "16A34A"},
        right={'val': 'single', 'sz': '4', 'color': HEX_BORDER}
    )

    pc0 = cb0.paragraphs[0]
    pc0.paragraph_format.space_before = Pt(0)
    pc0.paragraph_format.space_after = Pt(0)
    pc0.paragraph_format.line_spacing = 1.05
    rc0 = pc0.add_run(vuln_code)
    rc0.font.name = "Consolas"
    rc0.font.size = Pt(7.5)
    rc0.font.color.rgb = COLOR_SLATE

    pc1 = cb1.paragraphs[0]
    pc1.paragraph_format.space_before = Pt(0)
    pc1.paragraph_format.space_after = Pt(0)
    pc1.paragraph_format.line_spacing = 1.05
    rc1 = pc1.add_run(sec_code)
    rc1.font.name = "Consolas"
    rc1.font.size = Pt(7.5)
    rc1.font.color.rgb = COLOR_SLATE

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(3)

def add_image_box(doc, img_path, caption, width=Inches(6.2)):
    if not os.path.exists(img_path):
        print(f"Warning: image path does not exist: {img_path}")
        return

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(img_path, width=width)

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(5)
    rc = p_cap.add_run(f"Figure: {caption}")
    rc.font.name = "Calibri"
    rc.font.size = Pt(8.5)
    rc.font.italic = True
    rc.font.color.rgb = COLOR_MUTED

def add_dual_image_comparison(doc, img_vuln, img_sec, cap_vuln, cap_sec):
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(3.2)
    tbl.columns[1].width = Inches(3.2)

    c0, c1 = tbl.cell(0, 0), tbl.cell(0, 1)
    set_cell_margins(c0, top=30, bottom=30, left=30, right=30)
    set_cell_margins(c1, top=30, bottom=30, left=30, right=30)

    if os.path.exists(img_vuln):
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p0.add_run().add_picture(img_vuln, width=Inches(3.1))
        p0_cap = c0.add_paragraph()
        p0_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p0_cap.add_run(f"🔴 {cap_vuln}")
        r.font.name = "Calibri"
        r.font.size = Pt(8)
        r.font.bold = True
        r.font.color.rgb = COLOR_RED

    if os.path.exists(img_sec):
        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.add_run().add_picture(img_sec, width=Inches(3.1))
        p1_cap = c1.add_paragraph()
        p1_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p1_cap.add_run(f"🟢 {cap_sec}")
        r.font.name = "Calibri"
        r.font.size = Pt(8)
        r.font.bold = True
        r.font.color.rgb = COLOR_GREEN

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(3)

print("Writing document chapters...")
